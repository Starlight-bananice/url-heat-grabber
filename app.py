from __future__ import annotations

import contextlib
import os
import platform
import queue
import subprocess
import threading
import time
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook

import engine


APP_TITLE = '链接热度抓取工具'


def app_data_dir():
    system = platform.system()
    if system == 'Windows':
        roaming = os.environ.get('APPDATA')
        if roaming:
            return Path(roaming) / 'URLHeat'
        return Path.home() / 'AppData' / 'Roaming' / 'URLHeat'
    if system == 'Darwin':
        return Path.home() / 'Library' / 'Application Support' / 'URLHeat'
    data_home = os.environ.get('XDG_DATA_HOME')
    if data_home:
        return Path(data_home) / 'URLHeat'
    return Path.home() / '.local' / 'share' / 'URLHeat'


DATA_DIR = app_data_dir()
OUTPUT_DIR = Path.home() / 'Documents' / '链接热度抓取'


class QueueWriter:
    def __init__(self, events):
        self.events = events

    def write(self, value):
        value = value.strip()
        if value:
            self.events.put(('log', value))
        return len(value)

    def flush(self):
        return None


def read_links(path):
    path = Path(path)
    if path.suffix.lower() == '.xlsx':
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            values = [str(row[0]).strip() for row in worksheet.iter_rows(values_only=True) if row and row[0]]
        finally:
            workbook.close()
        return values
    return [line.strip() for line in path.read_text(encoding='utf-8-sig').splitlines() if line.strip()]


def write_private_file(path, content):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding='utf-8')
    try:
        path.chmod(0o600)
    except OSError:
        pass


class UrlHeatApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry('980x760')
        self.minsize(820, 620)
        self.protocol('WM_DELETE_WINDOW', self.close_app)

        self.events = queue.Queue()
        self.worker_thread = None
        self.running = False
        self.output_path = None

        self.input_path = tk.StringVar()
        self.output_path_var = tk.StringVar(value=str(OUTPUT_DIR))
        self.mode = tk.StringVar(value='1')
        self.captcha = tk.BooleanVar(value=False)
        self.resume = tk.BooleanVar(value=True)
        self.toutiao_workers = tk.IntVar(value=2)
        self.username = tk.StringVar()
        self.password = tk.StringVar()
        self.status = tk.StringVar(
            value='就绪；可拖动两条横线调整链接区和日志区大小。首次运行需要联网。'
        )

        self.build_ui()
        self.after(100, self.drain_events)

    def build_ui(self):
        root = ttk.Frame(self, padding=12)
        root.pack(fill='both', expand=True)
        root.columnconfigure(1, weight=1)
        root.rowconfigure(2, weight=1)

        ttk.Label(root, text='输入文件').grid(row=0, column=0, sticky='w', padx=(0, 8), pady=5)
        ttk.Entry(root, textvariable=self.input_path).grid(row=0, column=1, sticky='ew', pady=5)
        ttk.Button(root, text='选择并加载', command=self.choose_input).grid(row=0, column=2, padx=(8, 0), pady=5)

        ttk.Label(root, text='输出目录').grid(row=1, column=0, sticky='w', padx=(0, 8), pady=5)
        ttk.Entry(root, textvariable=self.output_path_var).grid(row=1, column=1, sticky='ew', pady=5)
        ttk.Button(root, text='选择目录', command=self.choose_output).grid(row=1, column=2, padx=(8, 0), pady=5)

        self.resize_panes = tk.PanedWindow(
            root,
            orient=tk.VERTICAL,
            borderwidth=0,
            relief=tk.FLAT,
            sashwidth=7,
            sashpad=1,
            sashrelief=tk.RAISED,
            showhandle=True,
            handlesize=12,
            opaqueresize=True,
        )
        self.resize_panes.grid(
            row=2, column=0, columnspan=3, sticky='nsew', pady=(8, 0)
        )

        input_frame = ttk.LabelFrame(
            self.resize_panes, text='链接（每行一个，也可以直接粘贴）', padding=8
        )
        input_frame.columnconfigure(0, weight=1)
        input_frame.rowconfigure(0, weight=1)
        self.input_text = tk.Text(input_frame, wrap='none', undo=True, height=12)
        self.input_text.grid(row=0, column=0, sticky='nsew')
        scrollbar = ttk.Scrollbar(input_frame, orient='vertical', command=self.input_text.yview)
        scrollbar.grid(row=0, column=1, sticky='ns')
        self.input_text.configure(yscrollcommand=scrollbar.set)

        middle_frame = ttk.Frame(self.resize_panes, padding=(0, 8, 0, 8))
        middle_frame.columnconfigure(0, weight=1)

        settings = ttk.LabelFrame(middle_frame, text='处理设置', padding=8)
        settings.grid(row=0, column=0, sticky='ew', pady=(0, 8))
        ttk.Radiobutton(settings, text='链接 + 互动数', variable=self.mode, value='1').grid(row=0, column=0, sticky='w')
        ttk.Radiobutton(settings, text='仅判断链接', variable=self.mode, value='0').grid(row=0, column=1, sticky='w', padx=(18, 0))
        ttk.Checkbutton(settings, text='验证码辅助', variable=self.captcha).grid(row=0, column=2, sticky='w', padx=(18, 0))
        ttk.Checkbutton(settings, text='从上次进度继续', variable=self.resume).grid(row=0, column=3, sticky='w', padx=(18, 0))
        ttk.Label(settings, text='头条并发').grid(row=1, column=0, sticky='w', pady=(8, 0))
        ttk.Spinbox(settings, from_=1, to=3, width=5, textvariable=self.toutiao_workers).grid(row=1, column=1, sticky='w', pady=(8, 0))
        ttk.Label(settings, text='抖音固定 1 个 worker，避免提高风控概率').grid(row=1, column=2, columnspan=2, sticky='w', padx=(18, 0), pady=(8, 0))

        credentials = ttk.LabelFrame(
            middle_frame, text='验证码账号（仅勾选“验证码辅助”时需要）', padding=8
        )
        credentials.grid(row=1, column=0, sticky='ew', pady=(0, 8))
        ttk.Label(credentials, text='账号').grid(row=0, column=0, sticky='w')
        ttk.Entry(credentials, textvariable=self.username, width=22).grid(row=0, column=1, sticky='w', padx=(6, 18))
        ttk.Label(credentials, text='密码').grid(row=0, column=2, sticky='w')
        ttk.Entry(credentials, textvariable=self.password, show='*', width=22).grid(row=0, column=3, sticky='w', padx=6)

        controls = ttk.Frame(middle_frame)
        controls.grid(row=2, column=0, sticky='ew')
        self.start_button = ttk.Button(controls, text='开始处理', command=self.start)
        self.start_button.pack(side='left')
        self.stop_button = ttk.Button(controls, text='停止（完成当前页后停止）', command=self.stop, state='disabled')
        self.stop_button.pack(side='left', padx=(8, 0))
        self.open_button = ttk.Button(controls, text='打开结果目录', command=self.open_output_dir)
        self.open_button.pack(side='left', padx=(8, 0))
        ttk.Label(controls, textvariable=self.status).pack(side='left', padx=(18, 0), fill='x', expand=True)

        self.progress = ttk.Progressbar(middle_frame, mode='indeterminate')
        self.progress.grid(row=3, column=0, sticky='ew', pady=(10, 0))

        log_frame = ttk.LabelFrame(self.resize_panes, text='运行日志', padding=6)
        log_frame.columnconfigure(0, weight=1)
        log_frame.rowconfigure(0, weight=1)
        self.log_text = tk.Text(log_frame, height=8, state='disabled', wrap='word')
        self.log_text.grid(row=0, column=0, sticky='nsew')
        log_scrollbar = ttk.Scrollbar(log_frame, orient='vertical', command=self.log_text.yview)
        log_scrollbar.grid(row=0, column=1, sticky='ns')
        self.log_text.configure(yscrollcommand=log_scrollbar.set)

        self.resize_panes.add(input_frame, minsize=180, stretch='always')
        self.resize_panes.add(middle_frame, minsize=245, stretch='always')
        self.resize_panes.add(log_frame, minsize=120, stretch='always')
        self.after_idle(self.set_default_pane_positions)

    def set_default_pane_positions(self):
        height = self.resize_panes.winfo_height()
        if height <= 0:
            return
        top_min = 180
        middle_min = 245
        log_min = 120
        if height < top_min + middle_min + log_min:
            return
        top = max(top_min, min(int(height * 0.45), height - middle_min - log_min))
        bottom = max(top + middle_min, height - log_min)
        self.resize_panes.sash_place(0, 0, top)
        self.resize_panes.sash_place(1, 0, bottom)

    def choose_input(self):
        path = filedialog.askopenfilename(
            title='选择链接文件',
            filetypes=[('链接文件', '*.txt *.csv *.xlsx'), ('所有文件', '*.*')],
        )
        if not path:
            return
        try:
            links = read_links(path)
            self.input_path.set(path)
            self.input_text.delete('1.0', 'end')
            self.input_text.insert('1.0', '\n'.join(links))
            self.status.set(f'已加载 {len(links)} 条链接。')
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f'读取文件失败：{type(exc).__name__}: {exc}')

    def choose_output(self):
        path = filedialog.askdirectory(title='选择结果目录', initialdir=self.output_path_var.get())
        if path:
            self.output_path_var.set(path)

    def open_output_dir(self):
        path = Path(self.output_path_var.get()).expanduser()
        path.mkdir(parents=True, exist_ok=True)
        if platform.system() == 'Windows':
            os.startfile(str(path))
        elif platform.system() == 'Darwin':
            subprocess.run(['open', str(path)], check=False)
        else:
            subprocess.run(['xdg-open', str(path)], check=False)

    def append_log(self, value):
        self.log_text.configure(state='normal')
        self.log_text.insert('end', value + '\n')
        self.log_text.see('end')
        self.log_text.configure(state='disabled')

    def start(self):
        if self.running:
            return
        links = [line.strip() for line in self.input_text.get('1.0', 'end').splitlines() if line.strip()]
        if not links:
            messagebox.showwarning(APP_TITLE, '请先导入或粘贴链接。')
            return
        try:
            workers = max(1, min(3, int(self.toutiao_workers.get())))
        except (TypeError, ValueError, tk.TclError):
            messagebox.showwarning(APP_TITLE, '头条并发必须是 1～3。')
            return

        data_dir = DATA_DIR
        output_dir = Path(self.output_path_var.get()).expanduser()
        data_dir.mkdir(parents=True, exist_ok=True)
        output_dir.mkdir(parents=True, exist_ok=True)
        write_private_file(data_dir / 'urls.txt', '\n'.join(links) + '\n')
        write_private_file(
            data_dir / 'settings.txt',
            f'只判断链接or抓取互动数和链接判断(0/1):{self.mode.get()}\n'
            f'是/否需要打码平台处理（需要耗费金额）(1/0):{1 if self.captcha.get() else 0}\n',
        )
        if self.captcha.get():
            write_private_file(data_dir / 'yzm.txt', f'账号:{self.username.get()}\n密码:{self.password.get()}\n')

        self.output_path = output_dir / f'链接判断结果_{time.strftime("%Y-%m-%d__%H-%M-%S")}.xlsx'
        args = [
            '--input', str(data_dir / 'urls.txt'),
            '--output', str(self.output_path),
            '--toutiao-workers', str(workers),
            '--douyin-workers', '1',
            '--other-workers', '1',
            '--checkpoint-every', '50',
        ]
        if self.resume.get():
            args.append('--resume')

        engine.OPT_BASE_DIR = data_dir
        engine.OPT_STOP_EVENT.clear()
        self.running = True
        self.start_button.configure(state='disabled')
        self.stop_button.configure(state='normal')
        self.progress.configure(mode='indeterminate')
        self.progress.start(12)
        self.status.set('正在处理；首次运行可能下载匹配的浏览器和 Driver。')
        self.append_log(f'开始处理 {len(links)} 条链接。')
        self.worker_thread = threading.Thread(target=self.run_engine, args=(args,), daemon=True)
        self.worker_thread.start()

    def run_engine(self, args):
        writer = QueueWriter(self.events)
        try:
            with contextlib.redirect_stdout(writer), contextlib.redirect_stderr(writer):
                return_code = engine.opt_main(args)
            self.events.put(('done', return_code))
        except Exception as exc:
            self.events.put(('fatal', f'{type(exc).__name__}: {exc}'))

    def stop(self):
        if self.running:
            engine.OPT_STOP_EVENT.set()
            self.status.set('已请求停止，将在当前页面处理结束后保存结果。')
            self.stop_button.configure(state='disabled')

    def drain_events(self):
        try:
            while True:
                event, value = self.events.get_nowait()
                if event == 'log':
                    self.append_log(value)
                elif event == 'done':
                    self.finish(value)
                elif event == 'fatal':
                    self.finish(1)
                    messagebox.showerror(APP_TITLE, value)
        except queue.Empty:
            pass
        self.after(100, self.drain_events)

    def finish(self, return_code):
        self.running = False
        self.progress.stop()
        self.start_button.configure(state='normal')
        self.stop_button.configure(state='disabled')
        if return_code == 0:
            self.status.set(f'处理完成，结果：{self.output_path}')
        else:
            self.status.set(f'已保存部分结果：{self.output_path}；可再次勾选“从上次进度继续”。')
        self.append_log(self.status.get())

    def close_app(self):
        if self.running:
            messagebox.showinfo(APP_TITLE, '请先点击停止，等待当前页面保存完成后再关闭。')
            return
        self.destroy()


if __name__ == '__main__':
    app = UrlHeatApp()
    app.mainloop()
