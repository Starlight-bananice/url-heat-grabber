import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import engine
from ui_model import (EXPORT_HEADERS, TaskStore, describe_result, metric_text, parse_links, platform_name,
                      read_links, write_result_workbook)


class InputTests(unittest.TestCase):
    def test_share_text_keeps_duplicate_rows_and_query_tokens(self):
        url = 'https://www.xiaohongshu.com/explore/123?xsec_token=A%2B%2F_b==&xsec_source=share'
        batch = parse_links(f'链接\n分享内容 {url}。\n{url}\n无效内容')
        self.assertEqual(batch.links, [url, url])
        self.assertEqual(batch.duplicates, 1)
        self.assertEqual(batch.ignored_lines, 1)
        self.assertEqual(batch.platforms, {'小红书': 2})
        self.assertEqual(parse_links(f'{url}\n{url}', True).links, [url])

    def test_reject_invalid_and_authenticated_urls(self):
        batch = parse_links('ftp://example.com\nhttps:///missing\nhttps://user:secret@example.com/a\nhttps://example.com:bad/a')
        self.assertEqual(batch.links, [])
        self.assertEqual(batch.ignored_lines, 4)

    def test_parentheses_and_query_punctuation_are_preserved(self):
        url = 'https://example.com/a_(test)?key=abc;'
        self.assertEqual(parse_links(url).links, [url])
        self.assertEqual(parse_links('[网页](https://example.com/a)').links, ['https://example.com/a'])

    def test_supported_domains_are_checked_as_hostnames(self):
        self.assertEqual(platform_name('https://weibo.com/123/456'), '微博')
        self.assertEqual(platform_name('https://news.sina.com.cn/abc'), '新浪')
        self.assertFalse(engine.opt_is_supported('https://weibo.com.example.org/123'))
        self.assertFalse(engine.opt_is_supported('https://example.org/?url=weibo.com'))
        self.assertEqual(engine.opt_group('https://example.org/?url=douyin.com'), 'unsupported')

    def test_excel_any_column_hyperlinks_and_formulas(self):
        with TemporaryDirectory() as root:
            path = Path(root) / 'links.xlsx'
            wb = Workbook()
            sheet = wb.active
            sheet.append(['序号', '标题', '链接'])
            sheet.append([1, '文字标题', 'https://weibo.com/123/456'])
            sheet.cell(3, 2, '点击查看').hyperlink = 'https://www.toutiao.com/article/123/'
            sheet.cell(4, 3, '=HYPERLINK("https://www.bilibili.com/video/BV123/", "查看")')
            wb.save(path)
            wb.close()
            batch = parse_links('\n'.join(read_links(path)))
            self.assertEqual(batch.links, ['https://weibo.com/123/456', 'https://www.toutiao.com/article/123/', 'https://www.bilibili.com/video/BV123/'])

    def test_csv_quoted_cells_bom_and_legacy_encoding(self):
        with TemporaryDirectory() as root:
            path = Path(root) / 'links.csv'
            path.write_text('标题,链接\n"一,二",https://weibo.com/123/456\n', encoding='utf-8-sig')
            self.assertEqual(parse_links('\n'.join(read_links(path))).links, ['https://weibo.com/123/456'])

    def test_csv_multiline_cell_preserves_separate_links(self):
        with TemporaryDirectory() as root:
            path = Path(root) / 'links.csv'
            path.write_text('标题,链接\n合并单元格,"https://weibo.com/1/1\nhttps://weibo.com/1/2"\n')
            self.assertEqual(parse_links('\n'.join(read_links(path))).links,
                             ['https://weibo.com/1/1', 'https://weibo.com/1/2'])
            path.write_text('标题,链接\n一二,https://weibo.com/123/456\n', encoding='gb18030')
            self.assertEqual(parse_links('\n'.join(read_links(path))).links, ['https://weibo.com/123/456'])


class ResultTests(unittest.TestCase):
    def test_zero_is_data_and_empty_is_unknown(self):
        self.assertEqual(describe_result({'点赞': 0}).label, '已获取数据')
        self.assertEqual(describe_result({'点赞': '0'}).label, '已获取数据')
        self.assertEqual(describe_result({'点赞': ''}).label, '暂无互动数据')
        self.assertEqual(describe_result({}, mode='0').label, '可访问')
        for zero in (0, '0', 0.0, ' 0 ', '0.00', '0万'):
            self.assertEqual(metric_text(zero), '—')
        self.assertEqual(metric_text(10), '10')
        self.assertEqual(metric_text('0.1万'), '0.1万')
        self.assertEqual(metric_text(None), '—')
        self.assertFalse(describe_result({'点赞': ''}).review)
        self.assertEqual(describe_result({'点赞': ''}).tone, 'muted')

    def test_error_status_takes_priority_over_metrics(self):
        for status in ('需验证', '访问受限', '处理失败', '不支持'):
            info = describe_result({'链接状态': status, '点赞': 2})
            self.assertEqual(info.label, status)
            self.assertTrue(info.review)
        self.assertFalse(describe_result({'链接状态': '未处理'}).review)

    def test_deleted_keeps_its_status_without_attention(self):
        for mode in ('0', '1'):
            info = describe_result({'链接状态': '已删除', '点赞': 2}, mode)
            self.assertEqual(info.label, '已删除')
            self.assertEqual(info.tone, 'muted')
            self.assertFalse(info.review)

    def test_export_pending_rows_zero_hyperlinks_and_safe_cells(self):
        with TemporaryDirectory() as root:
            path = Path(root) / 'result.xlsx'
            urls = ['https://weibo.com/1/1', 'https://www.bilibili.com/video/BV1/']
            results = {1: {'链接': urls[0], '点赞': 0, '评论/回复': '=1+1', '收藏': '0', '分享/转发': 0.0, '播放/阅读': 10}}
            write_result_workbook(path, results, urls)
            self.assertEqual(results[1]['点赞'], 0)
            wb = load_workbook(path, data_only=False)
            try:
                sheet = wb.active
                self.assertEqual(sheet.max_row, 3)
                self.assertEqual(tuple(cell.value for cell in sheet[1]), EXPORT_HEADERS)
                self.assertIsNone(sheet['C2'].value)
                self.assertIsNone(sheet['D2'].value)
                self.assertIsNone(sheet['F2'].value)
                self.assertIsNone(sheet['G2'].value)
                self.assertEqual(sheet['H2'].value, 10)
                self.assertEqual(sheet['E2'].data_type, 's')
                self.assertEqual(sheet['C3'].value, '未处理')
                self.assertEqual(sheet['B2'].hyperlink.target, urls[0])
                self.assertEqual(sheet.freeze_panes, 'C2')
                self.assertEqual(sheet.auto_filter.ref, 'A1:H3')
            finally:
                wb.close()

    def test_failed_export_does_not_replace_existing_result(self):
        with TemporaryDirectory() as root:
            path = Path(root) / 'result.xlsx'
            path.write_bytes(b'original')
            with patch('openpyxl.workbook.workbook.Workbook.save', side_effect=PermissionError('locked')):
                with self.assertRaises(PermissionError):
                    write_result_workbook(path, {}, [])
            self.assertEqual(path.read_bytes(), b'original')
            self.assertFalse(path.with_name('result.tmp.xlsx').exists())


class TaskFlowTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.base = patch.object(engine, 'OPT_BASE_DIR', self.root)
        self.base.start()
        engine.OPT_STOP_EVENT.clear()
        (self.root / 'settings.txt').write_text('mode:1\n')

    def tearDown(self):
        self.base.stop()
        self.temp.cleanup()
        engine.OPT_STOP_EVENT.clear()

    def prepare(self, urls):
        (self.root / 'urls.txt').write_text('\n'.join(urls))
        return ['--input', str(self.root / 'urls.txt'), '--output', str(self.root / 'result.xlsx')]

    def test_event_sequence_and_unknown_site_is_not_success(self):
        events = []
        args = self.prepare(['https://example.com/test'])
        with patch('engine.opt_create_driver') as driver:
            code = engine.opt_main(args, on_event=lambda kind, data: events.append((kind, data)))
        self.assertEqual(code, 0)
        driver.assert_not_called()
        self.assertEqual(events[0][0], 'prepared')
        result = next(data for kind, data in events if kind == 'result')
        self.assertEqual(result['row']['链接状态'], '不支持')
        self.assertEqual(events[-1][0], 'saved')
        self.assertTrue((self.root / 'result.xlsx').is_file())

    def test_resume_skips_only_final_results_and_reports_restored(self):
        urls = ['https://weibo.com/1/1', 'https://weibo.com/1/2', 'https://weibo.com/1/3']
        args = self.prepare(urls)
        checkpoint = {1: engine.opt_result_row(urls[0], metrics=('3', '', '', '', '')),
                      2: engine.opt_result_row(urls[1], status='需验证')}
        engine.opt_write_checkpoint(self.root / '链接判断结果_进行中.json', checkpoint, 3, engine.opt_signature(urls, '1'))
        processed, events = [], []
        def worker(bucket, group, driver_path, mode, system, config, callback):
            for index, url in bucket:
                processed.append(index)
                callback(index, engine.opt_result_row(url, metrics=('0', '', '', '', '')))
        with patch('engine.opt_worker', worker):
            engine.opt_main(args + ['--resume'], on_event=lambda kind, data: events.append((kind, data)))
        self.assertEqual(processed, [2, 3])
        self.assertEqual(list(events[0][1]['restored']), [1])
        self.assertEqual([data['completed'] for kind, data in events if kind == 'result'], [2, 3])

    def test_stop_requested_before_workers_is_not_lost(self):
        args = self.prepare(['https://weibo.com/1/1'])
        def stop_on_prepare(kind, data):
            if kind == 'prepared':
                engine.OPT_STOP_EVENT.set()
        with patch('engine.opt_create_driver') as driver:
            code = engine.opt_main(args, on_event=stop_on_prepare)
        driver.assert_not_called()
        self.assertEqual(code, 1)
        wb = load_workbook(self.root / 'result.xlsx')
        self.assertEqual(wb.active['C2'].value, '未处理')
        wb.close()

    def test_duplicates_fetch_once_and_export_every_original_row(self):
        a, b = 'https://weibo.com/1/1', 'https://weibo.com/1/2'
        urls = [a, b, a, b]
        args = self.prepare(urls)
        processed, events = [], []
        def worker(bucket, group, driver_path, mode, system, config, callback):
            for index, url in bucket:
                processed.append(index)
                callback(index, engine.opt_result_row(url, metrics=(str(index), '', '', '', '')))
        with patch('engine.opt_worker', worker):
            self.assertEqual(engine.opt_main(args, on_event=lambda kind, row: events.append((kind, row))), 0)
        self.assertEqual(processed, [1, 2])
        self.assertEqual([data['completed'] for kind, data in events if kind == 'result'], [1, 2, 3, 4])
        workbook = load_workbook(self.root / 'result.xlsx')
        try:
            sheet = workbook.active
            self.assertEqual(sheet.max_row, 5)
            self.assertEqual([sheet.cell(row, 2).value for row in range(2, 6)], urls)
            self.assertEqual([sheet.cell(row, 4).value for row in range(2, 6)], ['1', '2', '1', '2'])
        finally:
            workbook.close()

    def test_resume_reuses_a_later_duplicate_without_dropping_positions(self):
        a, b = 'https://weibo.com/1/1', 'https://weibo.com/1/2'
        urls = [a, b, a, b]
        args = self.prepare(urls)
        engine.opt_write_checkpoint(self.root / '链接判断结果_进行中.json', {3: engine.opt_result_row(a, metrics=('8', '', '', '', ''))}, 4, engine.opt_signature(urls, '1'))
        processed, events = [], []
        def worker(bucket, group, driver_path, mode, system, config, callback):
            for index, url in bucket:
                processed.append(index)
                callback(index, engine.opt_result_row(url, status='需验证'))
        with patch('engine.opt_worker', worker):
            self.assertEqual(engine.opt_main(args + ['--resume'], on_event=lambda kind, data: events.append((kind, data))), 0)
        self.assertEqual(processed, [2])
        self.assertEqual(set(events[0][1]['restored']), {1, 3})
        results = engine.opt_load_checkpoint(self.root / '链接判断结果_进行中.json', urls, engine.opt_signature(urls, '1'))
        self.assertEqual(set(results), {1, 3})

    def test_stop_keeps_duplicate_rows_and_pending_rows_in_export(self):
        a, b = 'https://weibo.com/1/1', 'https://weibo.com/1/2'
        args = self.prepare([a, b, a])
        def worker(bucket, group, driver_path, mode, system, config, callback):
            index, url = bucket[0]
            callback(index, engine.opt_result_row(url, metrics=('0', '', '', '', '')))
            engine.OPT_STOP_EVENT.set()
        with patch('engine.opt_worker', worker):
            self.assertEqual(engine.opt_main(args), 1)
        workbook = load_workbook(self.root / 'result.xlsx')
        try:
            self.assertEqual(workbook.active.max_row, 4)
            self.assertEqual([workbook.active.cell(row, 2).value for row in (2, 3, 4)], [a, b, a])
            self.assertEqual(workbook.active['C3'].value, '未处理')
            self.assertEqual(workbook.active['D2'].value, workbook.active['D4'].value)
        finally:
            workbook.close()

    def test_save_failure_has_no_saved_event(self):
        events = []
        args = self.prepare(['https://example.com/test'])
        with patch('engine.opt_save_xlsx', side_effect=PermissionError('locked')):
            with self.assertRaises(PermissionError):
                engine.opt_main(args, on_event=lambda kind, data: events.append(kind))
        self.assertNotIn('saved', events)
        self.assertTrue((self.root / '链接判断结果_进行中.json').is_file())

    def test_task_history_has_unique_outputs_and_mode_scoped_checkpoints(self):
        store = TaskStore(self.root)
        urls = ['https://weibo.com/1/1']
        sig = engine.opt_signature(urls, '1')
        first = store.create(urls, '1', sig, self.root)
        second = store.create(urls, '1', sig, self.root)
        self.assertNotEqual(first['output'], second['output'])
        checkpoint = store.directory(first) / '链接判断结果_进行中.json'
        engine.opt_write_checkpoint(checkpoint, {1: engine.opt_result_row(urls[0])}, 1, sig)
        self.assertEqual(store.latest_checkpoint(sig), checkpoint)
        self.assertIsNone(store.latest_checkpoint(engine.opt_signature(urls, '0')))
        self.assertEqual(len(store.records()), 2)
        self.assertEqual(store.records()[0]['id'], second['id'])
        self.assertEqual(len(store.results(first)), 1)


class TaskDeletionTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.store = TaskStore(self.root / 'data')
        self.record = self.store.create(['https://example.com/a'], '1', 'signature', self.root)
        self.output = Path(self.record['output'])
        self.output.write_bytes(b'task output')
        self.source = self.root / 'input.xlsx'
        self.source.write_bytes(b'original input')

    def tearDown(self):
        self.temp.cleanup()

    def test_delete_record_only_keeps_excel_and_original_input(self):
        result = self.store.delete_records([self.record])
        self.assertEqual(result['deleted'], [self.record['id']])
        self.assertTrue(self.output.is_file())
        self.assertTrue(self.source.is_file())
        self.assertEqual(self.store.records(), [])

    def test_delete_task_and_all_registered_exports(self):
        copy = self.root / '另存结果.xlsx'
        copy.write_bytes(b'export')
        self.store.register_export(self.record, copy)
        progress = self.store.directory(self.record) / '链接判断结果_进行中.json'
        progress.write_text('{}')
        result = self.store.delete_records([self.record], delete_files=True)
        self.assertEqual(result['errors'], {})
        self.assertFalse(self.output.exists())
        self.assertFalse(copy.exists())
        self.assertFalse(progress.exists())
        self.assertTrue(self.source.exists())
        self.assertIsNone(self.store.latest_checkpoint('signature'))

    def test_file_lock_preserves_task_for_retry(self):
        unlink = Path.unlink
        def locked(path, *args, **kwargs):
            if path == self.output:
                raise PermissionError('Excel is using this file')
            return unlink(path, *args, **kwargs)
        with patch.object(Path, 'unlink', locked):
            result = self.store.delete_records([self.record], delete_files=True)
        self.assertEqual(result['deleted'], [])
        self.assertIn(self.record['id'], result['errors'])
        self.assertTrue(self.output.exists())
        self.assertEqual(len(self.store.records()), 1)

    def test_shared_file_is_kept_until_both_records_are_selected(self):
        second = self.store.create(['https://example.com/b'], '1', 'other', self.root)
        self.store.register_export(second, self.output)
        result = self.store.delete_records([self.record], delete_files=True)
        self.assertTrue(self.output.exists())
        self.assertEqual(result['kept_files'], [str(self.output)])
        self.store.delete_records([second], delete_files=True)
        self.assertFalse(self.output.exists())

    def test_missing_output_does_not_block_history_cleanup(self):
        self.output.unlink()
        result = self.store.delete_records([self.record], delete_files=True)
        self.assertEqual(result['errors'], {})
        self.assertEqual(self.store.records(), [])

    def test_directory_traversal_and_linked_task_directory_are_rejected(self):
        with self.assertRaises(ValueError):
            self.store.delete_records([dict(self.record, id='../outside')], True)
        directory = self.store.directory(self.record)
        is_symlink = Path.is_symlink
        with patch.object(Path, 'is_symlink', lambda path: path == directory or is_symlink(path)):
            with self.assertRaises(ValueError):
                self.store.delete_records([self.record], True)
        self.assertTrue(self.output.exists())
        self.assertTrue(directory.is_dir())


if __name__ == '__main__':
    unittest.main()
