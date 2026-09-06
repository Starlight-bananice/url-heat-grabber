"""Opt-in packaging check, using an isolated workbook and a local HTML page.

Only the navigation destination is redirected. The packaged Tk UI, worker,
Selenium Manager, Chrome, parser, checkpoint and Excel exporter run normally.
This module is never executed during an ordinary application launch.
"""
import json
import faulthandler
import platform
import sys
import threading
import time
import traceback
from pathlib import Path

import selenium
from openpyxl import Workbook, load_workbook

import engine
from ui_model import EXPORT_HEADERS


def run_smoke_test(app_class, directory):
    directory = Path(directory).resolve()
    directory.mkdir(parents=True, exist_ok=True)
    report = {'ok': False, 'frozen': bool(getattr(sys, 'frozen', False)),
              'system': platform.system(), 'python': platform.python_version()}
    app = None
    trace = (directory / 'runtime-trace.txt').open('w', encoding='utf-8')
    faulthandler.enable(file=trace)
    faulthandler.dump_traceback_later(30, repeat=True, file=trace)
    original_navigate = engine.opt_navigate
    original_create = engine.opt_create_driver

    def save_report():
        temporary = directory / 'report.tmp'
        temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        temporary.replace(directory / 'report.json')

    try:
        url = 'https://tieba.baidu.com/p/123'
        values = {'share_pb': '7', 'comment_pb': '42', 'agree_pb': '108', 'collect': '6'}
        actions = ''.join(
            f'<div class="action-item"><svg><use href="#{icon}"></use></svg>'
            f'<span class="action-number">{value}</span></div>'
            for icon, value in values.items())
        fixture = directory / '帖子样本.html'
        fixture.write_text('<!doctype html><meta charset="utf-8"><title>本地测试帖子</title>'
                           f'<div class="pc-pb-first-floor-interactive">{actions}</div>', encoding='utf-8')
        workbook = Workbook()
        workbook.active.append(['链接', '备注'])
        for link in [url, url, 'https://example.com/unsupported']:
            workbook.active.append([link, '受控测试'])
        input_file = directory / '导入样本.xlsx'
        workbook.save(input_file)
        workbook.close()

        def navigate(driver, current_url, config):
            if current_url != url:
                raise AssertionError(f'Unexpected navigation: {current_url}')
            original_navigate(driver, fixture.as_uri(), config)
            report['navigations'] = report.get('navigations', 0) + 1

        def create_driver(*args, **kwargs):
            driver = original_create(*args, **kwargs)
            report['browser'] = driver.capabilities.get('browserVersion')
            report['driver'] = driver.capabilities.get('chrome', {}).get('chromedriverVersion')
            report['driver_path'] = driver.service.path
            return driver

        engine.opt_navigate = navigate
        engine.opt_create_driver = create_driver
        manager_name = 'selenium-manager.exe' if platform.system() == 'Windows' else 'selenium-manager'
        manager_os = 'windows' if platform.system() == 'Windows' else 'macos'
        manager = Path(selenium.__file__).resolve().parent / 'webdriver' / 'common' / manager_os / manager_name
        report['manager_bundled'] = manager.is_file()
        app = app_class(directory / '用户数据', directory / '导出结果')
        report['font'] = app.font_name
        report['tk'] = app.tk.call('info', 'patchlevel')
        app.import_file(input_file)
        assert app.refresh_input().links == [url, url, 'https://example.com/unsupported']
        started = time.monotonic()

        def fail(exc):
            report.update(error=f'{type(exc).__name__}: {exc}', logs=list(app.log_lines))
            save_report()
            engine.OPT_STOP_EVENT.set()
            app.destroy()

        def await_capture():
            if (directory / 'capture.done').exists() or time.monotonic() - started > 300:
                report['shutdown'] = 'requested'
                save_report()
                app.close_app()
                report['shutdown'] = 'window_closed'
                save_report()
            else:
                app.after(200, await_capture)

        def check_done():
            try:
                if app.running:
                    if time.monotonic() - started > 240:
                        raise TimeoutError('Packaged task did not finish in 240 seconds')
                    app.after(200, check_done)
                    return
                assert app.export_saved, app.status.get()
                assert [v.get() for v in app.stat_values] == ['3', '2', '1', '0']
                output = load_workbook(app.output_path, read_only=True, data_only=True)
                try:
                    rows = list(output.active.values)
                finally:
                    output.close()
                data = [dict(zip(rows[0], row)) for row in rows[1:]]
                assert rows[0] == EXPORT_HEADERS, rows[0]
                assert len(data) == 3, data
                expected = dict(zip(engine.OPT_RESULT_HEADERS[2:], ('108', '42', '6', '7', None)))
                assert {key: data[0][key] for key in expected} == expected, data
                assert {key: data[1][key] for key in expected} == expected, data
                assert [row['链接'] for row in data] == [url, url, 'https://example.com/unsupported']
                assert data[2]['链接状态'] == '不支持', data
                assert report['navigations'] == 1
                app.render_table()
                app.update_idletasks()
                assert len(app.tree.get_children()) == 3
                assert report.get('browser') and report['manager_bundled']
                report.update(ok=True, rows=data, stats=[v.get() for v in app.stat_values],
                              output=str(app.output_path), status=app.status.get(),
                              window=[app.winfo_width(), app.winfo_height()], logs=list(app.log_lines))
                save_report()
                app.after(200, await_capture)
            except Exception as exc:
                fail(exc)

        def start():
            try:
                app.start()
                app.after(200, check_done)
            except Exception as exc:
                fail(exc)

        app.report_callback_exception = lambda kind, exc, tb: fail(exc)
        app.after(600, start)
        app.mainloop()
        report['shutdown'] = 'mainloop_returned'
        report['threads'] = [{'name': thread.name, 'daemon': thread.daemon}
                             for thread in threading.enumerate()]
        save_report()
    except Exception:
        report['error'] = traceback.format_exc()
        save_report()
        if app is not None:
            app.destroy()
    finally:
        faulthandler.cancel_dump_traceback_later()
        faulthandler.disable()
        trace.close()
        engine.opt_navigate = original_navigate
        engine.opt_create_driver = original_create
