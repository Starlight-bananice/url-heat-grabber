"""Input, result presentation and local task storage for the desktop UI."""
from __future__ import annotations

import csv
import io
import json
import re
import shutil
import time
import uuid
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


METRICS = ('点赞', '评论/回复', '收藏', '分享/转发', '播放/阅读')
HEADERS = ('链接', '链接状态', *METRICS)
EXPORT_HEADERS = ('序号', *HEADERS)
PLATFORMS = (
    ('tieba.baidu.com', '百度贴吧'), ('mp.weixin.qq.com', '微信公众号'),
    ('douyin.com', '抖音'), ('iesdouyin.com', '抖音'), ('www.toutiao.com', '今日头条'),
    ('kuaishou.com', '快手'), ('weibo.com', '微博'), ('ixigua.com', '西瓜视频'),
    ('baidu.com', '百度'), ('163.com', '网易'), ('yoojia.com', '有驾'),
    ('uczzd.cn', 'UC'), ('mp.uc.cn', 'UC'), ('ifeng.com', '凤凰网'),
    ('sohu.com', '搜狐'), ('360kuai.com', '快资讯'), ('myzaker.com', 'ZAKER'),
    ('yidianzixun.com', '一点资讯'), ('html2.qktoutiao.com', '趣头条'),
    ('bilibili.com', 'B站'), ('dongchedi.com', '懂车帝'), ('news.qq.com', '腾讯新闻'),
    ('sina.com', '新浪'), ('sina.com.cn', '新浪'), ('sina.cn', '新浪'), ('iqiyi.com', '爱奇艺'),
    ('xiaohongshu.com', '小红书'),
)
URL_PATTERN = re.compile(r'''https?://[^\s<>"'\u201c\u201d\u300c\u300d]+''', re.I)
TRAILING_PUNCTUATION = '，。；！、）】》'
HEADER_WORDS = {'链接', '网址', 'url', 'urls', 'link', 'links', '标题', '平台', '序号'}


def platform_name(url):
    try:
        host = (urlsplit(url).hostname or '').lower().rstrip('.')
    except ValueError:
        return '不支持'
    for domain, name in PLATFORMS:
        if host == domain or host.endswith('.' + domain):
            return name
    return '不支持'


def valid_url(url):
    try:
        parts = urlsplit(url)
        return bool(parts.scheme.lower() in ('http', 'https') and parts.hostname
                    and '.' in parts.hostname and not parts.username and not parts.password
                    and parts.port != 0)
    except ValueError:
        return False


@dataclass
class LinkBatch:
    links: list[str]
    duplicates: int
    ignored_lines: int
    platforms: Counter

    @property
    def unsupported(self):
        return self.platforms.get('不支持', 0)


def parse_links(text, deduplicate=False):
    """Extract share URLs, preserving their query tokens and input order."""
    links, seen, duplicates, ignored = [], set(), 0, 0
    for line in text.splitlines():
        line = line.strip().lstrip('\ufeff')
        if not line:
            continue
        candidates = []
        for match in URL_PATTERN.findall(line):
            url = match.rstrip(TRAILING_PUNCTUATION)
            for opening, closing in (('(', ')'), ('[', ']'), ('{', '}')):
                while url.endswith(closing) and url.count(closing) > url.count(opening):
                    url = url[:-1]
            candidates.append(url)
        candidates = [url for url in candidates if valid_url(url)]
        if not candidates:
            if line.lower() not in HEADER_WORDS:
                ignored += 1
            continue
        for url in candidates:
            if url in seen:
                duplicates += 1
                if deduplicate:
                    continue
            seen.add(url)
            links.append(url)
    return LinkBatch(links, duplicates, ignored, Counter(map(platform_name, links)))


def read_links(path):
    """Read every cell, including Excel hyperlinks and quoted CSV fields."""
    path = Path(path)
    if path.suffix.lower() == '.xlsx':
        workbook = load_workbook(path, read_only=False, data_only=False)
        values = []
        try:
            for row in workbook.active.iter_rows():
                for cell in row:
                    target = cell.hyperlink.target if cell.hyperlink else None
                    if target:
                        values.append(target)
                    elif cell.value is not None:
                        values.append(str(cell.value))
        finally:
            workbook.close()
        return values
    try:
        text = path.read_text(encoding='utf-8-sig')
    except UnicodeDecodeError:
        text = path.read_text(encoding='gb18030')
    if path.suffix.lower() == '.csv':
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        return [cell for row in csv.reader(io.StringIO(text, newline=''), dialect) for cell in row if cell.strip()]
    return text.splitlines()


@dataclass(frozen=True)
class ResultInfo:
    label: str
    tone: str
    review: bool
    hint: str


def describe_result(row, mode='1'):
    status = str(row.get('链接状态') or '').strip()
    if status == '未处理':
        return ResultInfo('未处理', 'muted', False, '本次尚未处理；载入原任务后可继续。')
    if status == '已删除':
        return ResultInfo('已删除', 'muted', False, '页面明确提示内容不存在，或内容链接已跳回首页；可打开原网页复核。')
    explanations = {
        '不支持': ('muted', '当前解析器暂不支持此网址，未判断它是否有效；可双击打开原网页核实。'),
        '需验证': ('warning', '请在浏览器中核实页面；部分平台验证无法由自动抓取完成。'),
        '访问受限': ('warning', '平台限制了本次访问；建议稍后再试或手动核实，不代表链接已删除。'),
        '处理失败': ('danger', '本次浏览器或网络处理失败；请查看运行日志，稍后只重试这些链接。'),
    }
    if status in explanations:
        tone, hint = explanations[status]
        return ResultInfo(status, tone, True, hint)
    if status and status != '正常':
        return ResultInfo(status, 'warning', True, '页面返回了异常状态；建议双击打开原网页核实。')
    if mode == '0':
        return ResultInfo('可访问', 'success', False, '页面可访问；本次选择了仅检查链接，未抓取互动数。')
    if any(row.get(key) is not None and str(row.get(key)).strip() != '' for key in METRICS):
        return ResultInfo('已获取数据', 'success', False, '已获取页面公开互动数据；数值为 0 或未获取的指标显示为“—”，导出 Excel 时留空。')
    return ResultInfo('暂无互动数据', 'muted', False, '页面可访问，但未读取到公开互动数。')


def is_zero_metric(value):
    return bool(re.fullmatch(r'[+-]?0+(?:\.0+)?(?:万|亿)?', str(value).strip()))


def metric_text(value):
    return '—' if value is None or str(value).strip() == '' or is_zero_metric(value) else str(value)


def write_private_file(path, content):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + '.tmp')
    temporary.write_text(content, encoding='utf-8')
    try:
        temporary.chmod(0o600)
    except OSError:
        pass
    temporary.replace(path)


class TaskStore:
    def __init__(self, root):
        self.root = Path(root).resolve()

    def records(self):
        records = []
        if not (self.root / 'tasks').exists():
            return records
        for path in (self.root / 'tasks').glob('*/task.json'):
            try:
                record = json.loads(path.read_text(encoding='utf-8'))
                if (isinstance(record, dict) and record.get('id') == path.parent.name
                        and self.directory(record) == path.parent):
                    records.append(record)
            except (OSError, ValueError):
                continue
        return sorted(records, key=lambda item: (item.get('created_at', ''), item.get('created_ns', 0)), reverse=True)

    def create(self, links, mode, signature, output_dir):
        stamp = time.strftime('%Y-%m-%d %H:%M:%S')
        task_id = time.strftime('%Y%m%d-%H%M%S') + '-' + uuid.uuid4().hex[:8]
        record = dict(id=task_id, created_at=stamp, created_ns=time.time_ns(), links=links, mode=mode,
                      signature=signature, total=len(links), completed=0,
                      state='准备中', saved=False,
                      output=str(Path(output_dir).expanduser().absolute() / f'链接热度结果_{task_id}.xlsx'),
                      exports=[])
        self.save(record)
        return record

    def directory(self, record):
        task_id = record.get('id', '')
        if not isinstance(task_id, str) or not re.fullmatch(r'\d{8}-\d{6}-[0-9a-f]{8}', task_id):
            raise ValueError('任务编号无效')
        tasks = self.root / 'tasks'
        directory = tasks / task_id
        if any(path.is_symlink() or path.is_junction() for path in (tasks, directory)):
            raise ValueError('任务目录不能是符号链接')
        return directory

    def save(self, record):
        write_private_file(self.directory(record) / 'task.json', json.dumps(record, ensure_ascii=False, indent=2))

    def latest_checkpoint(self, signature):
        for record in self.records():
            path = self.directory(record) / '链接判断结果_进行中.json'
            if record.get('signature') == signature and path.is_file():
                return path
        return None

    def results(self, record):
        try:
            payload = json.loads((self.directory(record) / '链接判断结果_进行中.json').read_text(encoding='utf-8'))
            return {int(index): row for index, row in payload.get('results', {}).items() if isinstance(row, dict)}
        except (OSError, ValueError, TypeError):
            return {}

    @staticmethod
    def output_files(record):
        """Only explicit Excel outputs associated with this task; never scan folders."""
        paths = []
        exports = record.get('exports', [])
        if not isinstance(exports, list):
            raise ValueError('结果文件记录无效')
        for raw in [record.get('output'), *exports]:
            if not raw:
                continue
            if not isinstance(raw, str):
                raise ValueError('结果文件路径无效')
            path = Path(raw).expanduser().absolute()
            if path.suffix.lower() != '.xlsx':
                raise ValueError('仅可删除任务关联的 Excel 文件')
            if path not in paths:
                paths.append(path)
        return paths

    def register_export(self, record, path):
        exports = list(record.get('exports', []))
        value = str(Path(path).expanduser().absolute())
        if value not in exports:
            exports.append(value)
        record['exports'] = exports
        self.save(record)

    def deletion_files(self, records):
        selected = {record['id'] for record in records}
        retained = set()
        for other in self.records():
            if other['id'] not in selected:
                retained.update(path.resolve() for path in self.output_files(other))
        removable, shared = [], []
        seen = set()
        for record in records:
            self.directory(record)
            for path in self.output_files(record):
                if path in seen:
                    continue
                seen.add(path)
                (shared if path.resolve() in retained else removable).append(path)
        return removable, shared

    def delete_records(self, records, delete_files=False):
        """Delete selected task-owned data, retaining records on a deletion failure."""
        unique = {record['id']: record for record in records}
        directories = {key: self.directory(record) for key, record in unique.items()}
        removed_files, file_errors, kept_files = [], {}, []
        if delete_files:
            removable, shared = self.deletion_files(list(unique.values()))
            kept_files = [str(path) for path in shared]
            for path in removable:
                try:
                    existed = path.exists() or path.is_symlink()
                    path.unlink(missing_ok=True)
                    if existed:
                        removed_files.append(str(path))
                except OSError as exc:
                    file_errors[path] = str(exc)
        deleted, errors = [], {}
        for key, record in unique.items():
            failed = [path for path in self.output_files(record) if path in file_errors] if delete_files else []
            if failed:
                errors[key] = '结果文件删除失败：' + '；'.join(f'{path}：{file_errors[path]}' for path in failed)
                continue
            try:
                directory = directories[key]
                if directory.exists():
                    shutil.rmtree(directory)
                deleted.append(key)
            except OSError as exc:
                # Restore metadata if a partial directory removal reached task.json.
                errors[key] = f'任务记录删除失败：{exc}'
                try:
                    self.save(record)
                except (OSError, ValueError) as restore_error:
                    errors[key] += f'；记录恢复失败：{restore_error}'
        return dict(deleted=deleted, removed_files=removed_files, kept_files=kept_files, errors=errors)


def write_result_workbook(path, results, urls=None, mode='1'):
    """Export every input row, with explicit unprocessed states and useful formatting."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '抓取结果'
    sheet.append(EXPORT_HEADERS)
    indices = range(1, len(urls) + 1) if urls is not None else sorted(results)
    for index in indices:
        row = results.get(index, {'链接': urls[index - 1], '链接状态': '未处理'} if urls else {})
        row = dict(row)
        if urls is not None:
            row['链接'] = urls[index - 1]
        for metric in METRICS:
            if is_zero_metric(row.get(metric)):
                row[metric] = None
        values = (index, *(row.get(header, '') for header in HEADERS))
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            # Prevent spreadsheet formula interpretation of values supplied by a webpage.
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.font = Font(name='Microsoft YaHei', size=11, color='203832')
            cell.alignment = Alignment(vertical='center', wrap_text=False)
            if index % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F2F7F4')
        url_cell = sheet.cell(sheet.max_row, 2)
        if valid_url(str(url_cell.value or '')):
            url_cell.hyperlink = url_cell.value
            url_cell.font = Font(name='Microsoft YaHei', size=11, color='226B57', underline='single')
        sheet.row_dimensions[sheet.max_row].height = 25
    for cell in sheet[1]:
        cell.font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF', size=11)
        cell.fill = PatternFill('solid', fgColor='244B40')
        cell.alignment = Alignment(vertical='center')
    widths = (8, 65, 16, 14, 14, 14, 14, 16)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + '.tmp.xlsx')
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
